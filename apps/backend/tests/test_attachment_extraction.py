"""The parser is the only non-trivial branch in the attachment path — three
formats, a truncation rule, and a whitelist that is also a trust boundary."""

import io
import unittest

from app.services.attachments.extract import (
    MAX_EXTRACTED_CHARS,
    ExtractionError,
    extract_text,
)
from app.services.storage_service import kind_for_content_type
from app.utils.enums.attachment_enums import AttachmentKind

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _make_pptx(slides: list[tuple[str, str | None]]) -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    for title, notes in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _make_docx(paragraphs: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class KindClassificationTests(unittest.TestCase):
    def test_images_and_documents_are_separated(self):
        self.assertEqual(kind_for_content_type("image/png"), AttachmentKind.IMAGE)
        self.assertEqual(kind_for_content_type("application/pdf"), AttachmentKind.DOCUMENT)
        self.assertEqual(kind_for_content_type(PPTX), AttachmentKind.DOCUMENT)

    def test_charset_suffix_does_not_defeat_the_whitelist(self):
        self.assertEqual(
            kind_for_content_type("text/plain; charset=utf-8"), AttachmentKind.DOCUMENT
        )

    def test_anything_else_is_refused(self):
        # This is the trust boundary: an executable must not be storable by
        # claiming a content type the app never offers.
        for bad in ("application/zip", "application/x-msdownload", "", None):
            with self.assertRaises(ValueError):
                kind_for_content_type(bad)


class ExtractionTests(unittest.TestCase):
    def test_pptx_keeps_slide_numbers_and_speaker_notes(self):
        data = _make_pptx([("Revenue Growth", "Lead with the Q3 number")])
        text = extract_text(data, PPTX)
        # Slide numbers survive because "expand on slide 4" is a real revision
        # instruction, and notes are usually the most script-shaped text there is.
        self.assertIn("--- Slide 1 ---", text)
        self.assertIn("Revenue Growth", text)
        self.assertIn("Lead with the Q3 number", text)

    def test_docx_paragraphs_are_extracted(self):
        text = extract_text(_make_docx(["First point", "Second point"]), DOCX)
        self.assertIn("First point", text)
        self.assertIn("Second point", text)

    def test_plain_text_survives_a_bad_byte(self):
        text = extract_text(b"caf\xff nights", "text/plain")
        self.assertIn("nights", text)

    def test_long_documents_are_truncated_not_rejected(self):
        # Refusing here would cost the user their upload over a limit they have
        # no way to measure in advance.
        text = extract_text(("word " * 40_000).encode(), "text/plain")
        self.assertLessEqual(len(text), MAX_EXTRACTED_CHARS + 200)
        self.assertIn("document truncated", text)

    def test_a_document_with_no_text_is_an_error(self):
        with self.assertRaises(ExtractionError):
            extract_text(b"   \n  ", "text/plain")

    def test_unsupported_type_is_an_error(self):
        with self.assertRaises(ExtractionError):
            extract_text(b"anything", "application/zip")


class LinkWhitelistTests(unittest.TestCase):
    """The link field is a whitelist, so it gets the same treatment as the
    content-type one above."""

    @staticmethod
    def _links(values: list[str]) -> list[str]:
        from app.schemas.script_schema import ScriptGenerateRequest

        return ScriptGenerateRequest(
            description="a valid brief here",
            cardCount=5,
            durationMinutes=10,
            audience="general",
            links=values,
        ).links

    def test_bare_hostname_is_promoted_to_https(self):
        self.assertEqual(self._links(["example.com"]), ["https://example.com"])

    def test_http_and_https_pass_through(self):
        self.assertEqual(
            self._links(["http://a.io", "https://b.io/x"]),
            ["http://a.io", "https://b.io/x"],
        )

    def test_non_web_schemes_are_dropped(self):
        # The scheme has to be checked before the bare-hostname fixup: prefixing
        # first turns `javascript:alert(1)` into a URL that parses and passes.
        self.assertEqual(
            self._links(["javascript:alert(1)", "ftp://x.com", "data:text/html,x"]),
            [],
        )

    def test_blank_entries_are_dropped(self):
        self.assertEqual(self._links(["  ", ""]), [])


if __name__ == "__main__":
    unittest.main()
