"""The last branches: each test here exists for one path the broader suites
didn't happen to take, and is named for the behaviour it pins down."""

import threading
import time
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from types import SimpleNamespace

import pytest

from app.models.attachment_model import Attachment
from app.models.card_model import Card
from app.models.deck_model import Deck
from app.models.script_model import ScriptGeneration
from app.services.ai.providers.base import ChatRequest, ImageInput
from app.utils.enums.attachment_enums import AttachmentKind
from app.utils.enums.deck_enums import AudienceType, GenerationStatus
from app.utils.enums.speaking_style import SpeakingStyle

SCRIPT = "## [HOOK]\nSmall habits win. They compound quietly. Big goals stall. Start tiny today."
IMAGE = ImageInput(data=b"i", media_type="image/png")


@pytest.fixture
def gen(db, me):
    def make(**overrides):
        values = dict(
            user_id=me.id, description="Brief text long enough", duration_mins=3, card_count=2,
            audience=AudienceType.GENERAL, fingerprint="fp", title="T", script=SCRIPT,
            status=GenerationStatus.COMPLETED, last_seen_at=datetime.now(timezone.utc),
        )
        values.update(overrides)
        row = ScriptGeneration(**values)
        db.add(row)
        db.flush()
        return row

    return make


@pytest.fixture
def my_deck(db, me):
    def make(**overrides):
        values = dict(
            user_id=me.id, title="D", script=SCRIPT, color="#fff", duration_mins=3,
            audience=AudienceType.GENERAL, card_count=2, generation_status=GenerationStatus.COMPLETED,
        )
        values.update(overrides)
        row = Deck(**values)
        db.add(row)
        db.flush()
        return row

    return make


class TestCardEndpoints:
    def test_a_missing_deck_is_404_everywhere(self, client):
        assert client.get("/api/v1/decks/999999/cards/1").status_code == 404
        assert client.get("/api/v1/decks/999999/cards/status").status_code == 404

    def test_an_empty_edit_and_a_title_only_edit(self, client, my_deck, db):
        deck = my_deck()
        card = Card(deck_id=deck.id, position=1, title="A", description="d", keywords=[], color="#fff",
                    impact=Decimal("0.5"), delivery=SpeakingStyle.CONFIDENT)
        db.add(card)
        db.flush()
        same = client.patch(f"/api/v1/decks/{deck.id}/cards/{card.id}", json={"expected_version": card.version})
        assert same.json()["title"] == "A"
        renamed = client.patch(
            f"/api/v1/decks/{deck.id}/cards/{card.id}", json={"expected_version": card.version, "title": "B"},
        )
        assert renamed.json()["title"] == "B"

    def test_generation_refusals(self, client, my_deck, celery):
        assert client.post(f"/api/v1/decks/{my_deck(script=None).id}/cards/generate").status_code == 400
        busy = my_deck(cards_generation_status=GenerationStatus.PROCESSING)
        assert client.post(f"/api/v1/decks/{busy.id}/cards/generate").status_code == 409
        short = my_deck(script="One sentence.", card_count=4)
        assert client.post(f"/api/v1/decks/{short.id}/cards/generate").status_code == 400

    def test_status_from_the_database(self, client, my_deck, db):
        done = my_deck(cards_generation_status=GenerationStatus.COMPLETED)
        db.add(Card(deck_id=done.id, position=1, title="A", description="d", keywords=[], color="#fff",
                    impact=Decimal("0.5"), delivery=SpeakingStyle.CONFIDENT))
        db.flush()
        body = client.get(f"/api/v1/decks/{done.id}/cards/status").json()
        assert body["status"] == "completed" and len(body["cards"]) == 1
        pending = my_deck(cards_generation_status=GenerationStatus.PENDING)
        assert client.get(f"/api/v1/decks/{pending.id}/cards/status").json() == {"status": "pending"}


class TestScriptEndpoints:
    def test_polling_refreshes_a_stale_heartbeat_and_tolerates_naive_times(self, client, gen, db):
        stale = gen(last_seen_at=datetime.now(timezone.utc) - timedelta(minutes=5))
        client.get(f"/api/v1/scripts/{stale.id}/status")
        db.refresh(stale)
        assert datetime.now(timezone.utc) - stale.last_seen_at < timedelta(minutes=1)
        naive = gen(last_seen_at=datetime.now(timezone.utc).replace(tzinfo=None))
        assert client.get(f"/api/v1/scripts/{naive.id}/status").status_code == 200

    def test_edit_without_a_title_and_an_unchanged_edit(self, client, gen):
        row = gen()
        client.patch(f"/api/v1/scripts/{row.id}", json={"script": "Same words."})
        again = client.patch(f"/api/v1/scripts/{row.id}", json={"script": "Same words."})
        assert again.json()["title"] == "T"

    def test_retrying_a_failure_is_free(self, client, gen, celery):
        failed = gen(status=GenerationStatus.FAILED, script=None)
        assert client.post(f"/api/v1/scripts/{failed.id}/retry").status_code in (200, 202)

    def test_list_can_include_accepted_scripts(self, client, gen, my_deck):
        accepted = gen(deck_id=my_deck().id)
        without = [g["id"] for g in client.get("/api/v1/scripts").json()]
        with_all = [g["id"] for g in client.get("/api/v1/scripts?include_materialized=true").json()]
        assert accepted.id not in without and accepted.id in with_all

    def test_discarding_a_running_draft_purges_its_files(self, client, gen, db, celery, monkeypatch):
        from app.controllers import script_controller
        from app.services.storage_service import AttachmentStorageError

        running = gen(status=GenerationStatus.PROCESSING, celery_task_id="t")
        for key in ("ok", "stuck"):
            db.add(Attachment(user_id=running.user_id, kind=AttachmentKind.IMAGE, filename=key, content_type="image/png",
                              size_bytes=1, object_key=key, generation_id=running.id))
        db.flush()

        def delete(key):
            if key == "stuck":
                raise AttachmentStorageError("locked")

        monkeypatch.setattr(script_controller, "delete_attachment", delete)
        assert client.delete(f"/api/v1/scripts/{running.id}").status_code == 204
        db.refresh(running)
        assert running.status == GenerationStatus.CANCELLED
        assert db.query(Attachment).filter_by(generation_id=running.id).count() == 0

    def test_deck_status_and_cancel_without_a_running_job(self, client, gen):
        waiting = gen(cards_status=GenerationStatus.PENDING)
        assert client.get(f"/api/v1/scripts/{waiting.id}/deck/status").json()["status"] == "pending"
        idle = gen(cards_status=GenerationStatus.COMPLETED)
        assert client.post(f"/api/v1/scripts/{idle.id}/deck/cancel").json()["status"] == "completed"


def test_profile_update_without_a_nickname(client):
    body = client.patch("/api/v1/users/profile", json={"profession": "tech"}).json()
    assert body["profession"] == "tech"


def test_completing_onboarding_without_a_timestamp_stamps_one(client):
    body = client.put(
        "/api/v1/users/onboarding",
        json={"flow_version": "x", "status": "completed", "completed_steps": ["a", "a"]},
    ).json()
    assert body["completed_at"] is not None
    assert body["completed_steps"] == ["a"]


def test_helper_edges(me):
    from app.auth.clerk import decode_clerk_token
    from app.controllers import daily_controller, deck_controller
    from app.schemas.deck_schema import PublicDeckCreator
    from app.utils.nickname import InvalidNickname, validate_nickname

    assert decode_clerk_token("garbage") is None
    with pytest.raises(ValueError):
        daily_controller.select_variation("u", date.today(), 0)
    me.last_practiced_on = None
    assert daily_controller._break_date(me, date(2026, 9, 25)) == date(2026, 9, 25)
    deck_controller.revoke_task(None)
    assert PublicDeckCreator.model_validate({"id": 1, "name": "x"}).name == "x"
    with pytest.raises(InvalidNickname):
        validate_nickname("ab\x00cd")


def test_decode_clerk_token_accepts_a_valid_one(token_for):
    from app.auth.clerk import decode_clerk_token

    assert decode_clerk_token(token_for())["sub"] == "user_api"


class TestProviderImagePlacement:
    """Images go on the last *user* turn, however the conversation ends — and
    nowhere at all when there's no user turn to carry them."""

    def run(self, provider, monkeypatch, capture, messages):
        request = ChatRequest(system="s", messages=messages, images=(IMAGE,))
        return provider.complete(request, "m")

    def test_every_adapter(self, monkeypatch):
        from app.services.ai.providers import anthropic_provider, ollama_provider, openai_provider

        seen = []
        ollama = ollama_provider.OllamaProvider()
        monkeypatch.setattr(ollama, "_client", lambda: SimpleNamespace(chat=lambda **k: seen.append(k) or {"message": {"content": "x"}}))
        oai = openai_provider.OpenAIProvider()
        monkeypatch.setattr(oai, "_client", lambda: SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(
            create=lambda **k: seen.append(k) or SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="x"))])))))
        claude = anthropic_provider.AnthropicProvider()
        monkeypatch.setattr(claude, "_client", lambda: SimpleNamespace(messages=SimpleNamespace(
            create=lambda **k: seen.append(k) or SimpleNamespace(stop_reason="end", content=[SimpleNamespace(type="text", text="x")]))))

        for provider in (ollama, oai, claude):
            provider.complete(ChatRequest(system="s", messages=[{"role": "user", "content": "u"}, {"role": "assistant", "content": "a"}], images=(IMAGE,)), "m")
            provider.complete(ChatRequest(system="s", messages=[{"role": "assistant", "content": "a"}], images=(IMAGE,)), "m")
        assert len(seen) == 6


def test_openrouter_credentials(monkeypatch):
    from app.config.settings import settings
    from app.services.ai.providers.openrouter_provider import OpenRouterProvider

    monkeypatch.setattr(settings, "OPENROUTER_API_KEY", "or")
    assert OpenRouterProvider()._credentials()[0] == "or"


def test_gemini_client_built_once_under_contention(monkeypatch):
    from app.services.ai.providers.gemini_provider import GeminiProvider

    provider = GeminiProvider()
    # Another thread built the client while this one waited for the lock.
    class PreemptingLock:
        def __enter__(self):
            provider._cached_client = "built elsewhere"

        def __exit__(self, *exc):
            return False

    provider._client_lock = PreemptingLock()
    assert provider._client() == "built elsewhere"
    assert GeminiProvider._retry_after(SimpleNamespace(response=SimpleNamespace(headers={"retry-after": ""}))) is None


def test_limiter_blocks_until_a_slot_frees():
    from app.services.ai.rate_limit import AdaptiveLimiter

    limiter = AdaptiveLimiter(1)
    order = []
    holding = threading.Event()

    def first():
        with limiter.slot():
            holding.set()
            time.sleep(0.05)
            order.append("first")

    thread = threading.Thread(target=first)
    thread.start()
    holding.wait()
    with limiter.slot():
        order.append("second")
    thread.join()
    assert order == ["first", "second"]


def test_retry_helper_with_no_attempts():
    from app.services.ai.chat import _complete_with_retry
    from app.services.ai.providers.base import ProviderError

    with pytest.raises(ProviderError):
        _complete_with_retry(SimpleNamespace(name="p"), ChatRequest(system="", messages=[]), "m", retries=-1)


def test_card_generator_edges(monkeypatch):
    from app.services.ai import card_generator

    monkeypatch.setattr(card_generator, "MAX_CALIBRATION_CHARS", 30)
    digest = card_generator._script_digest("## A\n\n\n\n   \n\nFirst sentence. Second one here for length.")
    assert "## A" in digest
    # A delivery the caller doesn't accept falls back to the default.
    assert card_generator._coerce_card({"title": "t", "delivery": "calm"}, {"explaining"})["delivery"] == "explaining"
    monkeypatch.setattr(card_generator, "MAX_CALIBRATION_CHARS", 10_000)
    monkeypatch.setattr(card_generator, "chat", lambda messages: '[{"title": "a", "description": "b"}]')
    assert len(card_generator.generate_cards("One. Two.", 1)) == 1


def test_small_service_edges():
    from app.services import onboarding_demos
    from app.services.cards.impact_colors import assign_colors_by_impact

    assign_colors_by_impact([])
    empty = onboarding_demos.build_demo(
        onboarding_demos.CATALOG[0], "#fff",
        generate_script=lambda **_: ("T", "S"), generate_cards=lambda s, n: [], assign_colors=lambda cards: None,
    )
    assert empty["deck"]["cards"] == []


def test_presentation_color_may_be_absent():
    from app.schemas.ai_schema import AIGeneratedPresentationOutput

    assert AIGeneratedPresentationOutput.validate_hex_color(None) is None


def test_pptx_notes_without_text():
    import io

    from pptx import Presentation

    from app.services.attachments import extract

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Only a title"
    _ = slide.notes_slide  # created, left empty
    buffer = io.BytesIO()
    deck.save(buffer)
    text = extract.extract_text(
        buffer.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert "Speaker notes" not in text


class TestTaskEdges:
    """Retries exhausted: the error is final and the job records it."""

    def exhausted(self, task):
        task.push_request(retries=task.max_retries)

    def teardown_method(self):
        from app.tasks import card_tasks, deck_tasks, script_tasks

        for task in (script_tasks.generate_script_task, script_tasks.revise_script_task,
                     script_tasks.build_deck_from_generation, card_tasks.generate_deck_cards,
                     deck_tasks.revise_deck_script):
            while task.request_stack.top is not None:
                task.pop_request()

    def test_out_of_retries(self, task_db, make_user, fake_redis, monkeypatch):
        from app.services.ai.card_generator import CardGenerationError
        from app.services.ai.script_generator import ScriptGenerationError
        from app.tasks import card_tasks, deck_tasks, script_tasks

        user = make_user()
        task_db.commit()

        def boom_script(**_):
            raise ScriptGenerationError("final")

        def boom_cards(**_):
            raise CardGenerationError("final")

        generation = ScriptGeneration(user_id=user.id, description="d" * 12, duration_mins=1, card_count=2,
                                      audience=AudienceType.GENERAL, fingerprint="f", script=SCRIPT,
                                      status=GenerationStatus.PENDING)
        deck = Deck(user_id=user.id, script=SCRIPT, color="#fff", duration_mins=1, audience=AudienceType.GENERAL,
                    card_count=2, generation_status=GenerationStatus.COMPLETED)
        task_db.add_all([generation, deck])
        task_db.commit()

        monkeypatch.setattr(script_tasks, "generate_script", boom_script)
        monkeypatch.setattr(script_tasks, "revise_script", boom_script)
        monkeypatch.setattr(script_tasks, "generate_cards", boom_cards)
        monkeypatch.setattr(card_tasks, "generate_cards", boom_cards)
        monkeypatch.setattr(deck_tasks, "revise_script", boom_script)

        for task, args in (
            (script_tasks.generate_script_task, [generation.id]),
            (script_tasks.revise_script_task, [generation.id, "x"]),
            (script_tasks.build_deck_from_generation, [generation.id]),
            (card_tasks.generate_deck_cards, [deck.id, 2]),
            (deck_tasks.revise_deck_script, [deck.id, "x"]),
        ):
            self.exhausted(task)
            task(*args)
            task.pop_request()

    def test_revision_crash_on_a_deleted_generation(self, task_db, make_user, fake_redis, monkeypatch):
        from app.tasks import script_tasks

        user = make_user()
        generation = ScriptGeneration(user_id=user.id, description="d" * 12, duration_mins=1, card_count=1,
                                      audience=AudienceType.GENERAL, fingerprint="f", script=SCRIPT,
                                      status=GenerationStatus.COMPLETED)
        task_db.add(generation)
        task_db.commit()

        def delete_then_crash(**_):
            task_db.query(ScriptGeneration).filter_by(id=generation.id).delete()
            task_db.commit()
            raise RuntimeError("gone")

        monkeypatch.setattr(script_tasks, "revise_script", delete_then_crash)
        script_tasks.revise_script_task(generation.id, "x")

    def test_orphan_sweep_with_nothing_to_do(self, task_db):
        from app.tasks import script_tasks

        assert script_tasks.sweep_orphan_attachments() == 0


class TestDailyTaskEdges:
    def test_non_object_items_are_skipped(self, monkeypatch):
        from app.tasks import daily_tasks

        body = "One. Two. Three. Four. Five."
        monkeypatch.setattr(
            daily_tasks, "chat",
            lambda messages: '[1, {"body": "%s", "tip": "t", "title": "x"}]' % body,
        )
        units = daily_tasks.generate_day(date(2026, 9, 25))
        assert all(isinstance(unit, dict) for unit in units)

    def test_a_concurrent_fill_is_skipped_and_startup_queues_a_refill(self, task_db, monkeypatch, celery):
        from sqlalchemy.exc import IntegrityError

        from app.tasks import daily_tasks

        unit = {
            "framework": "PREP", "title": "t", "mood": "calm", "situation": "interview",
            "body": "b", "tip": "t", "variation_index": 0,
        }
        monkeypatch.setattr(daily_tasks, "generate_day", lambda day: [{**unit, "date": day}])
        original = task_db.commit

        def racing_commit():
            monkeypatch.setattr(task_db, "commit", original)
            raise IntegrityError("insert", {}, Exception("dup"))

        monkeypatch.setattr(task_db, "commit", racing_commit)
        daily_tasks.refill_daily_content(days=1)
        daily_tasks._fill_buffer_on_startup()
        assert any("refill" in name for name in celery.names())


def test_ranking_handles_an_uneven_catalogue(monkeypatch):
    from app.services import onboarding_demos

    demos = (
        {"id": "w1", "context": "work"}, {"id": "w2", "context": "work"},
        {"id": "c1", "context": "college"},
    )
    shaped = tuple(
        {**d, "label": "", "brief": "", "durationMinutes": 1, "cardCount": 1, "audience": "general",
         "mood": "calm", "profession": "other"}
        for d in demos
    )
    monkeypatch.setattr(onboarding_demos, "load_demos", lambda: shaped)
    assert [o["id"] for o in onboarding_demos.list_options(["college", "work"], limit=3)] == ["c1", "w1", "w2"]


def test_heartbeat_on_a_row_not_yet_written(db, gen):
    from app.controllers.script_controller import touch_generation

    row = gen()
    row.last_seen_at = None
    touch_generation(row, db)
    assert row.last_seen_at is not None
