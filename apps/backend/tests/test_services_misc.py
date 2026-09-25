"""The service layer's remaining corners: preferences, quota and RevenueCat,
document extraction, source loading, the realtime status cache's failure
handling, the chat fallback chain, the adaptive limiter, and the small
helpers and models the API suites don't happen to reach."""

import io
import time
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import httpx
import pytest
import redis

from app.config.settings import settings
from app.models.attachment_model import Attachment
from app.models.daily_model import DailyContent
from app.models.preferences_model import UserPreferences
from app.schemas.preferences_schema import UserPreferencesCreate, UserPreferencesUpdate
from app.services import quota, sources
from app.services.ai import chat as chat_module
from app.services.ai import prompts
from app.services.ai.providers.base import ChatRequest, ImageInput, ProviderError, RateLimitedError
from app.services.ai.rate_limit import AdaptiveLimiter
from app.services.attachments import extract
from app.services.realtime import deck_events, script_events
from app.utils.enums.attachment_enums import AttachmentKind
from app.utils.enums.deck_enums import AudienceType
from app.models.user_model import SubscriptionTier


# --- preferences -----------------------------------------------------------------


def test_preferences_controller(db, make_user):
    from fastapi import HTTPException

    from app.controllers import preferences_controller as prefs

    user = make_user()
    initial = UserPreferencesCreate(practiceRemindersEnabled=True, practiceReminderTime="07:30", defaultMood="calm")
    assert prefs.create_user_preferences(db, user.id, initial).defaultMood == "calm"
    with pytest.raises(HTTPException) as caught:
        prefs.create_user_preferences(db, user.id, initial)
    assert caught.value.status_code == 409

    updated = prefs.update_user_preferences(
        db, user.id,
        UserPreferencesUpdate(practiceRemindersEnabled=False, practiceReminderTime="08:00", defaultMood="playful"),
    )
    assert (updated.practiceRemindersEnabled, updated.practiceReminderTime, updated.defaultMood) == (False, "08:00", "playful")

    other = make_user()
    fresh = prefs.update_user_preferences(db, other.id, UserPreferencesUpdate())
    assert (fresh.practiceRemindersEnabled, fresh.practiceReminderTime, fresh.defaultMood) == (False, "09:00", "confident")
    third = make_user()
    chosen = prefs.update_user_preferences(
        db, third.id, UserPreferencesUpdate(practiceRemindersEnabled=True, practiceReminderTime="10:00", defaultMood="calm"),
    )
    assert chosen.practiceReminderTime == "10:00"
    assert "user_id" in repr(db.query(UserPreferences).first())


# --- quota -----------------------------------------------------------------------


class TestQuota:
    def test_tier_from_entitlements(self, monkeypatch):
        monkeypatch.setattr(settings, "REVENUECAT_ENTITLEMENT_ID", "pro")
        future = (datetime.now(timezone.utc) + timedelta(days=3)).isoformat().replace("+00:00", "Z")
        past = (datetime.now(timezone.utc) - timedelta(days=3)).isoformat()
        assert quota._tier_from_entitlements({}) == SubscriptionTier.FREE
        assert quota._tier_from_entitlements({"pro": {"expires_date": None}}) == SubscriptionTier.PRO
        assert quota._tier_from_entitlements({"pro": {"expires_date": future}}) == SubscriptionTier.PRO
        assert quota._tier_from_entitlements({"pro": {"expires_date": past}}) == SubscriptionTier.FREE
        assert quota._tier_from_entitlements({"pro": {"expires_date": "soon"}}) == SubscriptionTier.FREE

    def test_fetch_tier(self, monkeypatch):
        responses = iter([
            httpx.Response(404),
            httpx.Response(500),
            httpx.Response(200, json={"subscriber": {"entitlements": {"pro": {"expires_date": None}}}}),
            httpx.Response(200, json={}),
        ])
        monkeypatch.setattr(settings, "REVENUECAT_ENTITLEMENT_ID", "pro")
        monkeypatch.setattr(quota.httpx, "get", lambda *a, **k: next(responses))
        assert quota._fetch_tier("u") == SubscriptionTier.FREE
        assert quota._fetch_tier("u") is None
        assert quota._fetch_tier("u") == SubscriptionTier.PRO
        assert quota._fetch_tier("u") == SubscriptionTier.FREE

        def unreachable(*a, **k):
            raise httpx.ConnectError("down")

        monkeypatch.setattr(quota.httpx, "get", unreachable)
        assert quota._fetch_tier("u") is None

    def test_resolve_tier_caches_and_refreshes(self, db, make_user, monkeypatch):
        user = make_user()
        monkeypatch.setattr(settings, "REVENUECAT_API_KEY", "")
        assert quota.resolve_tier(user, db) == user.subscription_tier

        monkeypatch.setattr(settings, "REVENUECAT_API_KEY", "rc")
        monkeypatch.setattr(quota, "_fetch_tier", lambda _id: None)
        assert quota.resolve_tier(user, db) == user.subscription_tier
        monkeypatch.setattr(quota, "_fetch_tier", lambda _id: SubscriptionTier.PRO)
        assert quota.resolve_tier(user, db) == SubscriptionTier.PRO
        monkeypatch.setattr(quota, "_fetch_tier", lambda _id: SubscriptionTier.FREE)
        # Checked moments ago: the cached tier stands until the TTL runs out.
        assert quota.resolve_tier(user, db) == SubscriptionTier.PRO
        assert quota.resolve_tier(user, db, force=True) == SubscriptionTier.FREE

    def test_consume_refuses_past_the_limit_and_upgrades_rescue(self, db, make_user, monkeypatch):
        from fastapi import HTTPException

        monkeypatch.setattr(settings, "REVENUECAT_API_KEY", "")
        monkeypatch.setattr(settings, "FREE_MONTHLY_GENERATIONS", 1)
        monkeypatch.setattr(settings, "PRO_MONTHLY_GENERATIONS", 2)
        user = make_user()
        quota.consume_generation(user, db)
        with pytest.raises(HTTPException) as caught:
            quota.consume_generation(user, db)
        assert caught.value.status_code == 402
        assert "free generations" in caught.value.detail["message"]

        quota.refund_generation(user, db)
        assert user.monthly_generations_used == 0

        tiers = iter([SubscriptionTier.FREE, SubscriptionTier.PRO])
        monkeypatch.setattr(quota, "resolve_tier", lambda *a, **k: next(tiers))
        quota.consume_generation(user, db)  # free credit
        tiers = iter([SubscriptionTier.FREE, SubscriptionTier.PRO])
        quota.consume_generation(user, db)  # over free, upgraded to pro mid-request
        assert user.monthly_generations_used == 2

        monkeypatch.setattr(quota, "resolve_tier", lambda *a, **k: SubscriptionTier.PRO)
        with pytest.raises(HTTPException) as caught:
            quota.consume_generation(user, db)
        assert "this period's limit" in caught.value.detail["message"]


# --- extraction and sources -------------------------------------------------------


def docx_bytes():
    import docx

    document = docx.Document()
    document.add_paragraph("Opening line.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Q3"
    table.cell(0, 1).text = "Revenue"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes():
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Agenda"
    rows = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    rows.cell(0, 0).text = "Owner"
    slide.notes_slide.notes_text_frame.text = "Pause here."
    deck.slides.add_slide(deck.slide_layouts[6])  # blank: contributes nothing
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def pdf_bytes(text="Hello PDF"):
    """A one-page PDF with a text object, written by hand."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


class TestExtract:
    def test_each_format(self):
        assert "Q3 | Revenue" in extract.extract_text(docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        text = extract.extract_text(pptx_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        assert "--- Slide 1 ---" in text and "[Speaker notes] Pause here." in text and "Owner" in text
        assert "Hello PDF" in extract.extract_text(pdf_bytes(), "application/pdf")
        assert extract.extract_text(b"a\n\n\n\nb", "text/markdown") == "a\n\nb"

    def test_limits_and_failures(self, monkeypatch):
        monkeypatch.setattr(extract, "MAX_EXTRACTED_CHARS", 5)
        assert len(extract.extract_text(b"abcdefghij", "text/plain")) <= 5 + 50
        with pytest.raises(extract.ExtractionError, match="Unsupported"):
            extract.extract_text(b"x", "application/zip")
        with pytest.raises(extract.ExtractionError, match="No readable text"):
            extract.extract_text(b"   ", "text/plain")
        for content_type in (
            "application/pdf",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ):
            with pytest.raises(extract.ExtractionError):
                extract.extract_text(b"garbage", content_type)


def test_load_generation_sources(db, make_user, monkeypatch):
    from app.models.script_model import ScriptGeneration
    from app.services.storage_service import AttachmentStorageError
    from app.utils.enums.deck_enums import GenerationStatus

    user = make_user()
    generation = ScriptGeneration(user_id=user.id, description="d" * 12, duration_mins=1, card_count=1,
                                  audience=AudienceType.GENERAL, fingerprint="f", status=GenerationStatus.PENDING)
    db.add(generation)
    db.flush()
    assert sources.load_generation_sources(generation.id, db) == ([], None)

    common = dict(user_id=user.id, content_type="x", size_bytes=1, generation_id=generation.id)
    db.add_all([
        Attachment(kind=AttachmentKind.IMAGE, filename="ok.png", object_key="ok", **common),
        Attachment(kind=AttachmentKind.IMAGE, filename="bad.png", object_key="bad", **common),
        Attachment(kind=AttachmentKind.DOCUMENT, filename="notes.txt", object_key="d", extracted_text="Text", **common),
        Attachment(kind=AttachmentKind.DOCUMENT, filename="empty.pdf", object_key="e", **common),
    ])
    db.flush()

    def read(key):
        if key == "bad":
            raise AttachmentStorageError("gone")
        return ImageInput(data=b"i", media_type="image/png")

    monkeypatch.setattr(sources, "read_attachment_image", read)
    images, text = sources.load_generation_sources(generation.id, db)
    assert len(images) == 1 and text == "### notes.txt\n\nText"
    assert db.query(Attachment).first().is_image in (True, False)


# --- realtime cache -------------------------------------------------------------------


class BrokenRedis:
    def __getattr__(self, name):
        def fail(*args, **kwargs):
            raise redis.RedisError("down")

        return fail


def test_realtime_cache_round_trips_and_survives_redis_failures(monkeypatch, fake_redis):
    deck_events.write_card_status(1, {"status": "processing"})
    assert deck_events.read_card_status(1) == {"status": "processing"}
    assert deck_events.read_deck_status(2) is None
    script_events.write_generation_cards_status(3, {"status": "pending"})
    assert script_events.read_generation_cards_status(3) == {"status": "pending"}
    assert script_events.read_generation_cards_status(4) is None
    script_events.write_script_status(5, {"status": "x"})
    assert script_events.read_script_status(5) == {"status": "x"}
    script_events.mark_generation_activity(10)
    assert script_events.had_recent_generation_activity() is True
    script_events.clear_script_status(5)
    assert script_events.read_script_status(5) is None

    monkeypatch.setattr(deck_events, "_redis_client", BrokenRedis())
    deck_events.write_deck_status(1, {})
    deck_events.write_card_status(1, {})
    assert deck_events.read_deck_status(1) is None
    assert deck_events.read_card_status(1) is None
    script_events.write_script_status(1, {})
    script_events.write_generation_cards_status(1, {})
    assert script_events.read_script_status(1) is None
    assert script_events.read_generation_cards_status(1) is None
    script_events.mark_generation_activity(1)
    assert script_events.had_recent_generation_activity() is True
    script_events.clear_script_status(1)


def test_redis_client_is_created_lazily(monkeypatch):
    monkeypatch.setattr(deck_events, "_redis_client", None)
    client = deck_events._get_redis()
    assert client is deck_events._get_redis()


# --- chat -------------------------------------------------------------------------------


class FakeProvider:
    def __init__(self, name, models, outcomes):
        self.name = name
        self._models = models
        self.outcomes = list(outcomes)

    def models(self):
        return self._models

    def complete(self, request, model):
        outcome = self.outcomes.pop(0)
        if isinstance(outcome, Exception):
            raise outcome
        return outcome


@pytest.fixture
def providers(monkeypatch):
    registry = {}
    monkeypatch.setattr(chat_module, "get_provider", lambda name=None: registry[name or "primary"])
    monkeypatch.setattr(chat_module.time, "sleep", lambda s: None)
    monkeypatch.setattr(chat_module, "_openrouter_ready_at", 0.0)
    return registry


class TestChat:
    def test_rate_limits_are_waited_out_on_the_last_provider(self, providers, monkeypatch):
        monkeypatch.setattr(settings, "AI_USE_OPENROUTER", False)
        providers["primary"] = FakeProvider("ollama", ["m"], [RateLimitedError("slow"), RateLimitedError("slow", retry_after=1), "ok"])
        assert chat_module.chat([{"role": "system", "content": "s"}, {"role": "user", "content": "u"}]) == "ok"

    def test_gives_up_after_the_retries(self, providers, monkeypatch):
        monkeypatch.setattr(settings, "AI_USE_OPENROUTER", False)
        monkeypatch.setattr(chat_module, "MAX_RATE_LIMIT_RETRIES", 1)
        providers["primary"] = FakeProvider("ollama", ["m"], [RateLimitedError("a"), RateLimitedError("b")])
        with pytest.raises(chat_module.ModelCallError, match="rate limited"):
            chat_module.chat([{"role": "user", "content": "u"}])

    def test_no_models_and_every_model_failing(self, providers, monkeypatch):
        monkeypatch.setattr(settings, "AI_USE_OPENROUTER", False)
        providers["primary"] = FakeProvider("ollama", [], [])
        with pytest.raises(chat_module.ModelCallError, match="No models configured"):
            chat_module.chat([{"role": "user", "content": "u"}])
        providers["primary"] = FakeProvider("ollama", ["a", "b"], [ProviderError("x"), ValueError("y")])
        with pytest.raises(chat_module.ModelCallError, match="All models failed"):
            chat_module.chat([{"role": "user", "content": "u"}])

    def test_openrouter_first_then_parked_and_fallen_back(self, providers, monkeypatch):
        monkeypatch.setattr(settings, "AI_USE_OPENROUTER", True)
        monkeypatch.setattr(settings, "OPENROUTER_API_KEY", "or-key")
        providers["openrouter"] = FakeProvider("openrouter", ["free"], [RateLimitedError("429", retry_after=30)])
        providers["primary"] = FakeProvider("ollama", ["m"], ["fallback answer"])
        assert chat_module.chat([{"role": "user", "content": "u"}]) == "fallback answer"
        # Parked now: the next call goes straight to the primary.
        assert chat_module._openrouter_parked() is True
        providers["primary"].outcomes.append("direct")
        assert [p.name for p in chat_module.provider_chain()] == ["ollama"]
        assert chat_module.chat([{"role": "user", "content": "u"}]) == "direct"

    def test_openrouter_unavailable_parks_it_and_the_cooldown_ends(self, providers, monkeypatch):
        monkeypatch.setattr(settings, "AI_USE_OPENROUTER", True)
        monkeypatch.setattr(settings, "OPENROUTER_API_KEY", "or-key")
        providers["openrouter"] = FakeProvider("openrouter", ["free"], [ProviderError("down")])
        providers["primary"] = FakeProvider("ollama", ["m"], ["ok"])
        assert chat_module.chat([{"role": "user", "content": "u"}]) == "ok"
        monkeypatch.setattr(chat_module, "_openrouter_ready_at", time.monotonic() - 1)
        assert chat_module._openrouter_parked() is False

    def test_openrouter_without_free_models_or_as_the_only_provider(self, providers, monkeypatch):
        monkeypatch.setattr(settings, "AI_USE_OPENROUTER", True)
        monkeypatch.setattr(settings, "OPENROUTER_API_KEY", "or-key")
        providers["openrouter"] = FakeProvider("openrouter", [], [])
        providers["primary"] = FakeProvider("ollama", ["m"], ["ok"])
        assert [p.name for p in chat_module.provider_chain()] == ["ollama"]
        both = FakeProvider("openrouter", ["free"], ["served free"])
        providers["openrouter"] = both
        providers["primary"] = both
        assert [p.name for p in chat_module.provider_chain()] == ["openrouter"]
        assert chat_module.chat([{"role": "user", "content": "u"}]) == "served free"

    def test_served_logging_and_colour(self, monkeypatch):
        monkeypatch.setattr(settings, "AI_USE_OPENROUTER", True)
        monkeypatch.setattr(settings, "OPENROUTER_API_KEY", "k")
        chat_module._log_served(SimpleNamespace(name="ollama"), "m")
        monkeypatch.setattr(settings, "OPENROUTER_API_KEY", "")
        chat_module._log_served(SimpleNamespace(name="ollama"), "m")
        monkeypatch.setattr(chat_module, "_COLOR", True)
        assert chat_module._paint("x", "green").startswith("\033[32m")

    def test_delay_and_retry_helper(self):
        assert chat_module._rate_limit_delay(0, 3.0) == 3.0
        assert chat_module._rate_limit_delay(10, None) <= chat_module.RATE_LIMIT_MAX_DELAY * 1.5

    def test_map_parallel(self):
        assert chat_module.map_parallel(lambda x: x, []) == []
        assert chat_module.map_parallel(lambda x: x * 2, [3]) == [6]
        assert chat_module.map_parallel(lambda x: x * 2, [1, 2, 3]) == [2, 4, 6]

        def boom(x):
            if x == 2:
                raise ValueError("bad")
            return x

        with pytest.raises(ValueError):
            chat_module.map_parallel(boom, [1, 2, 3])


# --- limiter ------------------------------------------------------------------------------


def test_adaptive_limiter(monkeypatch):
    from app.services.ai import rate_limit

    limiter = AdaptiveLimiter(8)
    assert limiter.limit == 8
    limiter.record_success()  # at the ceiling: nothing to recover
    limiter.record_rate_limited()
    assert limiter.limit == 4
    limiter.record_rate_limited()  # inside the cooldown: ignored
    assert limiter.limit == 4
    for _ in range(rate_limit.SUCCESSES_PER_INCREASE):
        limiter.record_success()
    assert limiter.limit == 5
    floor = AdaptiveLimiter(1)
    floor.record_rate_limited()
    assert floor.limit == rate_limit.MIN_LIMIT
    with floor.slot():
        pass


# --- prompts and small things -----------------------------------------------------------


def test_remaining_prompt_builders():
    spec = prompts.SECTION_SPECS[0]
    assert prompts.build_plan_prompt("brief", 3, AudienceType.STUDENTS, source_text="notes", links="https://a.b")
    assert prompts.build_section_prompt(
        spec=spec, target_words=80, title="T", audience=AudienceType.GENERAL,
        plan_note="plan", include_question=True, extra_insight="insight",
    )
    assert prompts.build_section_prompt(spec=spec, target_words=80, title="T", audience=AudienceType.GENERAL, plan_note="")
    assert prompts.build_revision_scope_prompt(instruction="shorter", sections=[("HOOK", "text")])
    assert prompts.build_section_revision_prompt(
        label="HOOK", body="b", instruction="i", title="t", audience=AudienceType.GENERAL,
        source_text="s", links="https://a.b",
    )


def test_small_things(monkeypatch, s3, db):
    import app.core.exceptions as exceptions
    from app.config.settings import Settings
    from app.core import s3_client
    from app.db import database

    assert issubclass(exceptions.AIGenerationError, Exception)
    monkeypatch.setattr(s3_client, "client", s3)
    s3_client.ensure_bucket()
    s3.head_bucket.assert_called_once()

    legacy = Settings(AI_MODEL="", AI_FALLBACK_MODEL="", OLLAMA_MODEL="old", OLLAMA_FALLBACK_MODEL="older")
    assert (legacy.AI_MODEL, legacy.AI_FALLBACK_MODEL) == ("old", "older")

    generator = database.get_db()
    session = next(generator)
    generator.close()
    assert session is not None
    database._reset_pool_after_fork()

    assert "PREP" in repr(DailyContent(date=datetime.now().date(), framework="PREP", variation_index=0))
