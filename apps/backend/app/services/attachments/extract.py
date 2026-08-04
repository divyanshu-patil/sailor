"""Reduce an uploaded document to plain text.

Text, not vision. A slide deck rendered to images would carry its diagrams, but
it would also cost a vision-capable model and a large multiple of the tokens,
and this pipeline has to work identically on a local Ollama text model. Text
extraction is the format every provider accepts.

Extraction happens once, at upload, and the result is stored on the attachment
row — a revision months later is grounded in the same source material without
the original file needing to survive or be re-parsed.
"""

import io
import logging
import re

logger = logging.getLogger("celery")


class ExtractionError(Exception):
    """The document could not be read as text."""


# What the upload endpoint accepts, mapped to the reader that handles it. A
# content type absent from here is rejected at the boundary rather than stored
# and discovered to be unreadable by a worker later.
PDF_TYPES = {"application/pdf"}
DOCX_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
PPTX_TYPES = {
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
TEXT_TYPES = {"text/plain", "text/markdown"}

DOCUMENT_TYPES = PDF_TYPES | DOCX_TYPES | PPTX_TYPES | TEXT_TYPES

IMAGE_TYPES = {"image/jpeg", "image/jpg", "image/png"}

# A prompt has a budget and the script brief has to survive alongside the source
# material. ~40k characters is roughly 10k tokens — a long report or a 60-slide
# deck — and past that the tail is truncated rather than the upload refused: a
# user who attached a 300-page PDF still gets a script, grounded in its opening.
MAX_EXTRACTED_CHARS = 40_000

# Collapses the run of blank lines that every one of these formats produces on
# an empty paragraph or an empty table cell.
_BLANK_RUN = re.compile(r"\n{3,}")


def _clean(parts: list[str]) -> str:
    text = "\n".join(part.strip() for part in parts if part and part.strip())
    return _BLANK_RUN.sub("\n\n", text).strip()


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on install extras
        raise ExtractionError("PDF support requires the `pypdf` package.") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        return _clean([page.extract_text() or "" for page in reader.pages])
    except Exception as exc:
        raise ExtractionError("This PDF could not be read.") from exc


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - depends on install extras
        raise ExtractionError("Word support requires the `python-docx` package.") from exc

    try:
        document = docx.Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        # Tables hold a lot of what a document actually says — figures, comparisons,
        # timelines — and none of it appears in `paragraphs`.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return _clean(parts)
    except Exception as exc:
        raise ExtractionError("This Word document could not be read.") from exc


def _extract_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - depends on install extras
        raise ExtractionError(
            "PowerPoint support requires the `python-pptx` package."
        ) from exc

    try:
        presentation = Presentation(io.BytesIO(data))
        parts: list[str] = []
        for number, slide in enumerate(presentation.slides, start=1):
            # Slide numbers are kept deliberately. "Expand on slide 4" is a
            # revision instruction users actually give, and it only means
            # anything if the source text says which slide is which.
            slide_parts = [f"--- Slide {number} ---"]
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    slide_parts.append(shape.text_frame.text)
                elif getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            slide_parts.append(" | ".join(cells))
            # Speaker notes are the presenter's own words about the slide —
            # usually the most useful text in the file for writing a script.
            notes = slide.notes_slide if slide.has_notes_slide else None
            if notes is not None and notes.notes_text_frame is not None:
                note_text = notes.notes_text_frame.text.strip()
                if note_text:
                    slide_parts.append(f"[Speaker notes] {note_text}")
            if len(slide_parts) > 1:
                parts.append("\n".join(slide_parts))
        return _clean(parts)
    except Exception as exc:
        raise ExtractionError("This presentation could not be read.") from exc


def _extract_plain(data: bytes) -> str:
    # errors="replace" rather than raising: a stray byte in an otherwise fine
    # text file shouldn't cost the user their upload.
    return _clean([data.decode("utf-8", errors="replace")])


def extract_text(data: bytes, content_type: str) -> str:
    """Plain text from one document. Raises ExtractionError if unreadable.

    Truncation is silent by design — see MAX_EXTRACTED_CHARS. The alternative,
    refusing a long document, loses the user their upload over a limit they have
    no way to measure in advance.
    """
    content_type = (content_type or "").split(";")[0].strip().lower()

    if content_type in PDF_TYPES:
        text = _extract_pdf(data)
    elif content_type in DOCX_TYPES:
        text = _extract_docx(data)
    elif content_type in PPTX_TYPES:
        text = _extract_pptx(data)
    elif content_type in TEXT_TYPES:
        text = _extract_plain(data)
    else:
        raise ExtractionError(f"Unsupported document type: {content_type}")

    if not text:
        raise ExtractionError(
            "No readable text was found — a scanned or image-only document "
            "can't be used as source material."
        )

    if len(text) > MAX_EXTRACTED_CHARS:
        logger.info(f"[attachments] truncating extracted text {len(text)} -> {MAX_EXTRACTED_CHARS}")
        text = text[:MAX_EXTRACTED_CHARS].rsplit("\n", 1)[0] + "\n\n[... document truncated]"

    return text
